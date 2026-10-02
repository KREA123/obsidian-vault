# -*- coding: utf-8 -*-
"""Fișe de post (4), declarație de disponibilitate, politica de egalitate de șanse.
Rulare: python3 build_echipa.py <OUTDIR>. Cifrele vin din 03-NUCLEU §9 (echipa) și §13 (orizontale)."""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docxlib import write_runs, W
from data import TITLU, APPLICANT, ACRONIM

OUTDIR = sys.argv[1]
NAVY = RGBColor(0x1B, 0x3A, 0x6B)
APEL = 'PR BI P1/1.1/1/2026 — Sprijin pentru dezvoltarea unui model conceptual inovativ (Proof of Concept)'


def base_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(3)
    st.element.rPr.rFonts.set(W('eastAsia'), 'Calibri')
    for lvl, size in (('Heading 1', 13), ('Heading 2', 11)):
        h = d.styles[lvl]
        h.font.name = 'Calibri'; h.font.size = Pt(size); h.font.color.rgb = NAVY; h.font.bold = True
        h.paragraph_format.space_before = Pt(6); h.paragraph_format.space_after = Pt(2)
        rpr = h.element.get_or_add_rPr()
        rf = rpr.find(W('rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            rf.set(W(a), 'Calibri')
        for a in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            if rf.get(W(a)) is not None:
                del rf.attrib[W(a)]
    return d


def P(d, text, align=None, size=None, italic=False, bullet=False):
    p = d.add_paragraph(style='List Bullet' if bullet else None)
    write_runs(p._p, text, plain=False, size=int(size * 2) if size else None)
    if italic:
        for r in p.runs:
            r.italic = True
    p.alignment = {'c': WD_ALIGN_PARAGRAPH.CENTER, 'r': WD_ALIGN_PARAGRAPH.RIGHT,
                   'j': WD_ALIGN_PARAGRAPH.JUSTIFY}.get(align, p.alignment)
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(W('val'), 'clear'); sh.set(W('color'), 'auto'); sh.set(W('fill'), fill)
    tcpr.append(sh)


def table(d, header, data, widths, head=True):
    t = d.add_table(rows=1 if head else 0, cols=len(widths))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if head:
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]
            write_runs(c.paragraphs[0]._p, f'**{h}**', plain=False, color='FFFFFF', size=19)
            shade(c, '1B3A6B')
    for row in data:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            lines = v.split('\n')
            write_runs(cells[j].paragraphs[0]._p, lines[0], plain=False, size=19)
            for ln in lines[1:]:
                write_runs(cells[j].add_paragraph()._p, ln, plain=False, size=19)
    t.autofit = False
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    d.add_paragraph()
    return t


def note_box(d, text, fill='FFF4E5'):
    t = d.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    c = t.rows[0].cells[0]
    shade(c, fill)
    lines = text.split('\n')
    write_runs(c.paragraphs[0]._p, lines[0], plain=False, size=19)
    for ln in lines[1:]:
        write_runs(c.add_paragraph()._p, ln, plain=False, size=19)
    d.add_paragraph()


def header_block(d, title):
    P(d, f'**{APPLICANT}** · CUI [[CUI]] · [[J40/…/…]] · sediul: [[adresa, sector, București]]', size=9)
    P(d, f'Proiect: **{ACRONIM}** — {TITLU}', size=9)
    P(d, f'Apel: {APEL}', size=9)
    P(d, f'**{title}**', 'c', 14)


# --------------------------------------------------------------------------- FIȘE DE POST
COMMON_DUTIES = [
    'completează lunar **fișa de pontaj** (ore pe activități A1–A8) și un scurt **raport lunar de activitate**, semnate și predate managerului de proiect până în ziua 5 a lunii următoare;',
    'respectă **confidențialitatea** informațiilor tehnice și comerciale ale proiectului; nu publică rezultate înainte de depunerea cererilor de protecție IP („întâi protecția, apoi publicarea”);',
    f'recunoaște că **drepturile asupra rezultatelor** obținute în proiect aparțin {APPLICANT}, potrivit contractului de muncă și legii;',
    'declară în scris, în **3 zile lucrătoare**, orice situație de **conflict de interese** (inclusiv legături cu furnizorii proiectului, până la gradul II de rudenie);',
    'respectă principiile de **egalitate de șanse și nediscriminare** și regulile de **comunicare și vizibilitate** ale PR BI 2021–2027;',
    'nu desfășoară, în orele pontate în proiect, activități finanțate din alte surse publice (**fără dublă finanțare a timpului de lucru**).',
]

POSTS = [
    dict(
        file='Fisa_post_Manager_proiect',
        rol='Manager de proiect',
        titular='[[Nume Prenume — Andu]]',
        cor='[[Ct: cod COR, ex. „manager proiect”]]',
        activitate='Management de proiect (activitate conexă II.3) — A8',
        hg='Anexa 2 la HG 1188/2022, categoria 1 „director/responsabil/manager de program/proiect” — plafon 50 €/oră brut',
        tarif='18 €/oră brut (sub plafon), plătit în lei la cursul InforEuro din contract',
        timp='20 ore/lună × 24 luni = 480 ore (L1–L24); CIM cu timp parțial [[Ct: confirmă că administratorul poate avea CIM cu ARTEMIS; altfel managementul se face fără cost în proiect]]',
        sub='raportează asociatului unic/AGA ARTEMIS; coordonează întreaga echipă; interfața cu AM PR BI (ADR BI)',
        cerinte=[
            'studii superioare [[domeniul]]; experiență de management de proiecte și de echipe (KREA, MundiShop) — [[ani]];',
            'experiență de cercetare-dezvoltare documentată: inițiatorul și dezvoltatorul conceptului SOUL (firmware, serviciul AI, prototipul SoulOS, CAD — dovezile D1–D8 din Raportul științific);',
            'cunoașterea regulilor de achiziții pentru beneficiari privați (Ordinul MFE 1284/2016) sau angajamentul de a le însuși în prima lună;',
            'competențe digitale: MySMIS2021, semnătură electronică calificată.',
        ],
        atributii=[
            'planifică, coordonează și monitorizează activitățile A1–A8 și calendarul (Gantt L1–L24); convoacă ședința lunară a echipei;',
            'derulează **achizițiile** proiectului conform Ordinului MFE 1284/2016 (caiete de sarcini, publicare, evaluarea ofertelor, contracte), cu cerințele DNSH (ecodesign, RoHS) incluse;',
            'întocmește și transmite **rapoartele de progres**, cererile de prefinanțare/plată/rambursare și **raportările indicatorilor** (RCO 01, RCO 02, RCR 03; IE1 la L6, IE2 la L14, IE3 la L22);',
            'asigură relația cu AM PR BI: notificări în 5 zile lucrătoare, clarificări, vizite de monitorizare; păstrează **dosarul proiectului** (arhivă fizică și electronică);',
            'coordonează **protecția IP** (căutare de anteriorități, cererea de desen/model UE ≤ L5, marca UE ≤ L6) și **valorificarea**: customer discovery, transformarea scrisorilor de intenție în ≥ 2 acorduri, planul de comercializare (R3, R5);',
            'urmărește bugetul, contabilitatea analitică separată, pontajele și auditul financiar (II.4); implementează măsurile de comunicare și vizibilitate (II.5);',
            'aplică politica de egalitate de șanse și măsura suplimentară (panelul TRL 5, testul de accesibilitate); gestionează riscurile din Planul de afaceri §6.4.',
        ],
    ),
    dict(
        file='Fisa_post_Coordonator_tehnic',
        rol='Coordonator tehnic/științific',
        titular='[[Nume Prenume]]',
        cor='[[Ct: cod COR, ex. „cercetător științific” / „inginer electronist”]]',
        activitate='Activități CDI (I.3) — A1, A2, A3, A4, A5, A7',
        hg='Anexa 2 la HG 1188/2022, categoria 2 „responsabil tehnic proiect / personal de specialitate” — plafon 35 €/oră brut [[cat. 1 = 50 €/h dacă titularul este conferențiar/CS II]]',
        tarif='34 €/oră brut (sub plafon), plătit în lei',
        timp='50 ore/lună × 24 luni = 1.200 ore (L1–L24); CIM cu timp parțial (cumul permis cu postul de bază) **sau** detașare de la o organizație de cercetare',
        sub='în subordinea managerului de proiect pe partea administrativă; conduce echipa tehnică (inginer embedded, student)',
        cerinte=[
            '**titlul de doctor** în științe inginerești în domeniul proiectului (electronică, telecomunicații, calculatoare, sisteme embedded) **SAU** **≥ 5 ani** de experiență dovedită în domeniu **SAU** **≥ 2 proiecte CDI** conduse cu succes ca director/coordonator (Grila ETF 3.1a-i);',
            'experiență în proiectarea și validarea sistemelor embedded/IoT (microcontrolere ESP32/ARM, interfețe de afișare, radio BLE/Wi-Fi, măsurători de laborator);',
            'publicații sau rezultate CDI în domeniu (de preferat indexate WoS/Scopus) — necesare pentru coordonarea articolului open-access (R4);',
            'capacitatea de a redacta rapoarte tehnice de validare TRL (Anexa 18).',
        ],
        atributii=[
            'elaborează **specificația sistemului SOUL M** (L1.1, L2) și **protocoalele experimentale P1–P11** cu praguri de acceptare TRL 4/TRL 5 (L1.2, L3);',
            'aprobă bancul de test (M1, L3): aparatele de măsură, calibrarea, procedurile;',
            'coordonează realizarea modelului de laborator integrat (A2) și **validarea TRL 4** (A3); redactează **Raportul de validare TRL 4** (R1, L14 = IE2) și propune decizia de poartă go/no-go;',
            'coordonează tehnic prototipul pre-serie (A4): arhitectura PCB, antena prin fereastra polimerică, carcasa din aluminiu, securitatea EN 18031;',
            'coordonează **validarea TRL 5 în mediu relevant** (A5) și pre-testarea EMC/RF; redactează **Raportul de validare TRL 5** (R2, L22 = IE3);',
            'coordonează redactarea și trimiterea **articolului științific open-access** (R4) și prezentarea la conferință (A7); îndrumă studentul în activitatea de cercetare;',
            'semnalează managerului de proiect, cel târziu la L14 (≤ 2/3 din durată), orice risc de **eșec de cercetare** și întocmește, dacă e cazul, raportul de eșec (Anexa 17).',
        ],
    ),
    dict(
        file='Fisa_post_Inginer_embedded',
        rol='Inginer embedded/firmware (expert tehnic/cercetător)',
        titular='[[Nume Prenume]]',
        cor='[[Ct: cod COR, ex. „inginer electronist” / „inginer de sistem în informatică”]]',
        activitate='Activități CDI (I.3/I.4) — A2, A3, A4, A5',
        hg='Anexa 2 la HG 1188/2022, categoria 2 „personal de specialitate” — plafon 35 €/oră brut',
        tarif='25 €/oră brut (sub plafon), plătit în lei',
        timp='80 ore/lună × 21 luni = 1.680 ore (L2–L22); CIM cu timp parțial',
        sub='în subordinea coordonatorului tehnic/științific',
        cerinte=[
            '**studii superioare** în domeniul proiectului (electronică, calculatoare, automatică, telecomunicații) **și ≥ 2 ani** de experiență în domeniu **SAU ≥ 1 publicație** indexată Web of Science/Scopus în domeniu (Grila ETF 3.1a-ii);',
            'C/C++ pentru microcontrolere (ESP-IDF/Arduino/PlatformIO), drivere de afișaj și touch (RGB/SPI/I2C), audio I2S, BLE;',
            'proiectare PCB (KiCad/Altium) și punere în funcțiune; lucru cu osciloscop, analizor de consum;',
            'teste automate, control al versiunilor (git), documentare tehnică.',
        ],
        atributii=[
            'portează și optimizează firmware-ul SoulOS pe platforma SOUL M (afișaj 2,8″ ST7701, touch GT911, audio I2S, PMU, baterie) — modelul de laborator integrat (L2.1, IE1 la L6);',
            'implementează fluxul vocal și integrarea cu serviciul AI din UE; menține testele automate;',
            'execută protocoalele de laborator P1–P8 (fps, touch, tastare, latență, autonomie, offline) și raportează datele brute către coordonator (A3);',
            'proiectează, pune în funcțiune și testează **PCB-ul propriu** rev A/B și integrarea în carcasa din aluminiu (L4.1, A4); implementează secure boot și OTA semnat (EN 18031);',
            'asamblează, verifică (QC) și pregătește cele 15 prototipuri pre-serie (M3, L17); suport tehnic în validarea TRL 5 și la pre-testarea EMC/RF (A5);',
            'contribuie la rapoartele TRL 4/TRL 5 și la articolul științific (secțiunile tehnice).',
        ],
    ),
    dict(
        file='Fisa_post_Student_cercetare',
        rol='Student — asistent de cercetare',
        titular='[[Nume Prenume]]',
        cor='[[Ct: cod COR, ex. „asistent de cercetare”]]',
        activitate='Activități CDI (I.3) — A3, A5, A7',
        hg='Anexa 2 la HG 1188/2022, categoria 4 „activități suport / student” — plafon 15 €/oră brut',
        tarif='12 €/oră brut (sub plafon), plătit în lei',
        timp='60 ore/lună × 20 luni = 1.200 ore (L5–L24); CIM cu timp parțial, program flexibil compatibil cu cursurile',
        sub='în subordinea coordonatorului tehnic/științific',
        cerinte=[
            '**vârsta ≤ 24 de ani** la data depunerii cererii de finanțare și **înmatriculat** într-o instituție de învățământ superior în anul universitar curent (Grila ETF 4.3a) — se dovedesc cu copia actului de identitate și adeverința de student;',
            'student la o specializare tehnică (electronică, calculatoare, automatică, telecomunicații) sau de interacțiune om–calculator/științe cognitive;',
            'abilități de lucru cu date (Python sau echivalent), redactare în limba română și engleză;',
            'interes pentru testarea cu utilizatori și accesibilitate.',
        ],
        atributii=[
            'pregătește și asistă studiul de tastare (N = 6) și testele de laborator TRL 4 (A3): instrucțiuni pentru participanți, colectarea datelor, curățarea și analiza statistică;',
            'participă la **recrutarea panelului TRL 5** (≥ 40 % femei, ≥ 2 persoane cu deficiențe de vedere/motorii) și la organizarea celor 2 workshop-uri de testare (A5);',
            'administrează chestionarele (SUS, jurnale de utilizare), anonimizează setul de date și pregătește tabelele și graficele pentru raportul TRL 5;',
            'contribuie la construirea corpusului de test în limba română (comenzi vocale, tastare) și la documentația tehnică;',
            'contribuie la redactarea articolului open-access (metodă, rezultate) și la prezentarea de la conferință (A7).',
        ],
    ),
]


def build_post(p):
    d = base_doc()
    header_block(d, f'FIȘA POSTULUI — {p["rol"].upper()}')
    P(d, f'Nr. [[…]] / [[ZZ.LL.AAAA]] · anexă la [[CIM nr. … / declarația de disponibilitate]]', 'c', 9, italic=True)
    table(d, [], [
        ['**Denumirea postului în proiect**', p['rol']],
        ['**Titularul postului**', p['titular'] + '\n[[sau: POST VACANT — se ocupă după semnarea contractului de finanțare, pe baza cerințelor de mai jos]]'],
        ['**Cod COR**', p['cor']],
        ['**Activitatea / încadrarea în buget**', p['activitate']],
        ['**Categoria salarială**', p['hg']],
        ['**Tarif bugetat**', p['tarif']],
        ['**Timp de lucru și durată**', p['timp']],
        ['**Locul de muncă**', 'sediul/punctul de lucru ARTEMIS din București [[adresa]] și laboratoarele partenere; muncă hibridă permisă pentru analiză și redactare'],
        ['**Relații ierarhice și funcționale**', p['sub']],
    ], [5.0, 11.6], head=False)
    d.add_heading('1. Cerințele postului', level=2)
    for c in p['cerinte']:
        P(d, c, 'j', bullet=True)
    d.add_heading('2. Atribuții și responsabilități în proiect', level=2)
    for a in p['atributii']:
        P(d, a, 'j', bullet=True)
    d.add_heading('3. Obligații comune ale membrilor echipei', level=2)
    for a in COMMON_DUTIES:
        P(d, a, 'j', bullet=True)
    d.add_heading('4. Evaluare', level=2)
    P(d, 'Activitatea se evaluează pe baza livrabilelor și a termenelor din Cererea de finanțare și din Planul de monitorizare (Anexa 13), a pontajelor și a rapoartelor lunare.', 'j')
    d.add_paragraph()
    table(d, ['ANGAJATOR', 'TITULAR'], [
        [f'{APPLICANT}\nprin [[Nume Prenume]], administrator\n\nSemnătura: ____________________\nData: [[ZZ.LL.AAAA]]',
         'Am luat la cunoștință și îmi asum prezenta fișă a postului.\n[[Nume Prenume]]\n\nSemnătura: ____________________\nData: [[ZZ.LL.AAAA]]'],
    ], [8.3, 8.3])
    d.save(f'{OUTDIR}/{p["file"]}.docx')


for p in POSTS:
    build_post(p)

# --------------------------------------------------------------------------- DECLARAȚIE DE DISPONIBILITATE
d = base_doc()
P(d, '**MODEL — DECLARAȚIE DE DISPONIBILITATE ȘI ANGAJAMENT**', 'c', 14)
P(d, 'pentru personalul-cheie care nu este încă angajat al solicitantului (Grila ETF 3.1 și 3.2b: „contracte/angajamente/declarații de disponibilitate”)', 'c', 9, italic=True)
note_box(d, '**Cum se folosește (se șterge înainte de semnare):**\n'
         '1. O declarație separată pentru **fiecare** persoană (coordonatorul tehnic, inginerul embedded; opțional studentul).\n'
         '2. Se semnează **înainte de depunere**, olograf (apoi scanată integral) sau cu semnătură electronică calificată; se depune împreună cu **CV-ul Europass semnat**, **fișa postului semnată** și copiile diplomelor/certificărilor.\n'
         '3. Nimeni nu este angajat și nu este plătit înainte de semnarea contractului de finanțare. Angajarea (CIM cu timp parțial) se face după contract.\n'
         '4. Rolul, orele și perioada trebuie să fie **identice** cu Cererea de finanțare, Planul de afaceri §5.2 și fișa postului.')
P(d, '**DECLARAȚIE DE DISPONIBILITATE**', 'c', 12)
P(d, 'Subsemnatul/Subsemnata [[Nume Prenume]], posesor/posesoare al/a actului de identitate seria [[..]] nr. [[……]], cu domiciliul în [[……]], '
     f'în prezent [[funcția și angajatorul actual, ex. „șef de lucrări, Universitatea Națională de Știință și Tehnologie POLITEHNICA București, Facultatea ETTI”]],', 'j')
P(d, '**declar pe propria răspundere că:**', 'j')
for t in [
    f'sunt de acord să fac parte din echipa proiectului **{ACRONIM}** — „{TITLU}”, propus de **{APPLICANT}** în apelul {APEL}, în poziția de [[Coordonator tehnic/științific | Inginer embedded/firmware]];',
    'sunt disponibil(ă) să desfășor activitățile descrise în fișa postului anexată, cu o normă de [[50 | 80]] ore/lună, în perioada [[L1–L24 | L2–L22]] a proiectului, începând cu data semnării contractului de finanțare (estimat [[2027]]);',
    'accept să fiu angajat(ă) de solicitant cu **contract individual de muncă cu timp parțial** [[sau: detașat(ă) de la organizația de cercetare la care sunt angajat(ă)]], cu tariful orar brut prevăzut în bugetul proiectului;',
    'activitatea în proiect nu se suprapune cu programul meu de lucru de bază și, cumulat cu celelalte activități, respectă limitele legale ale timpului de muncă; nu există clauze contractuale (ex. neconcurență, exclusivitate) care să mă împiedice să particip [[sau: am acordul scris al angajatorului actual]];',
    'orele lucrate în proiect **nu vor fi finanțate din alte surse publice** (fără dublă finanțare);',
    'nu mă aflu în situație de **conflict de interese** față de solicitant, de reprezentantul său legal sau de evaluatorii apelului și voi declara orice astfel de situație în 3 zile lucrătoare;',
    'informațiile din **CV-ul meu** (anexat, semnat) și documentele care îl însoțesc sunt reale; cunosc prevederile art. 326 din Codul penal privind falsul în declarații;',
    'sunt de acord cu prelucrarea datelor mele personale de către solicitant și de către autoritățile programului, exclusiv în scopul evaluării, contractării, implementării și controlului proiectului (Regulamentul (UE) 2016/679).',
]:
    P(d, t, 'j', bullet=True)
P(d, '**Anexez:** CV Europass semnat · fișa postului semnată · copii ale diplomelor/certificărilor · [[lista publicațiilor cu DOI / dovezi ale proiectelor CDI]] · [[adeverință de vechime în domeniu]].', 'j')
d.add_paragraph()
table(d, [], [['Nume Prenume: [[……]]\nSemnătura: ____________________\nData: [[ZZ.LL.AAAA]]\nE-mail / telefon: [[……]]']], [16.6], head=False)
d.save(f'{OUTDIR}/Declaratie_disponibilitate_MODEL.docx')

# --------------------------------------------------------------------------- POLITICA DE EGALITATE DE ȘANSE
d = base_doc()
P(d, f'**{APPLICANT}** · CUI [[CUI]] · [[J40/…/…]] · [[adresa sediului]]', size=9)
P(d, '**DECIZIA ADMINISTRATORULUI** nr. [[…]] din [[ZZ.LL.AAAA]]', 'c', 13)
P(d, '**privind adoptarea Politicii interne de egalitate de șanse, nediscriminare și accesibilitate**', 'c', 11)
note_box(d, '**Se datează înainte de depunerea cererii de finanțare** și se anexează la cerere (Grila ETF 4.3b). Data și numărul se trec și în registrul de decizii al societății.')
P(d, f'Subsemnatul [[Nume Prenume]], în calitate de administrator al {APPLICANT}, având în vedere Ordonanța Guvernului nr. 137/2000 privind prevenirea și sancționarea tuturor formelor de discriminare, Legea nr. 202/2002 privind egalitatea de șanse și de tratament între femei și bărbați, Legea nr. 448/2006 privind protecția și promovarea drepturilor persoanelor cu handicap, art. 9 din Regulamentul (UE) 2021/1060 și Carta drepturilor fundamentale a Uniunii Europene,', 'j')
P(d, '**DECID:**', 'c')
P(d, '**Art. 1.** Se adoptă Politica internă de egalitate de șanse, nediscriminare și accesibilitate, prevăzută mai jos, aplicabilă tuturor angajaților, colaboratorilor și candidaților societății, inclusiv în proiectul ' + ACRONIM + '.', 'j')
P(d, '**Art. 2.** Responsabil cu aplicarea și monitorizarea politicii este [[Nume Prenume, funcția]]. Politica se aduce la cunoștința personalului prin semnătură și se publică pe site-ul societății.', 'j')
P(d, '**Art. 3.** Prezenta decizie intră în vigoare la data semnării.', 'j')
d.add_heading('POLITICA INTERNĂ DE EGALITATE DE ȘANSE, NEDISCRIMINARE ȘI ACCESIBILITATE', level=1)
sections = [
    ('1. Principii', [
        'Societatea interzice orice discriminare directă sau indirectă pe criterii de sex, origine rasială sau etnică, religie sau convingeri, dizabilitate, vârstă, orientare sexuală, stare civilă sau responsabilități familiale.',
        'Hărțuirea și hărțuirea sexuală sunt interzise; orice persoană poate sesiza confidențial responsabilul desemnat, fără risc de represalii.',
    ]),
    ('2. Recrutare și selecție', [
        'anunțurile de angajare sunt formulate neutru din punctul de vedere al genului și sunt publicate pe cel puțin două canale, inclusiv canale universitare;',
        'selecția se face pe criterii obiective, publicate în fișa postului (studii, experiență, competențe); întrebările despre starea civilă, sarcină sau planuri familiale sunt interzise;',
        'candidații cu dizabilități beneficiază, la cerere, de adaptări rezonabile ale procesului de selecție.',
    ]),
    ('3. Salarizare și condiții de muncă', [
        'salarizarea se face pe o grilă transparentă (tarif orar pe categorii de post), cu același tarif pentru aceeași muncă, indiferent de sex;',
        'program flexibil și muncă hibridă pentru activitățile care o permit; programul studenților se adaptează orarului academic;',
        'acces egal la formare profesională, conferințe și co-autorat științific, pe baza contribuției.',
    ]),
    ('4. Măsuri în proiectul ' + ACRONIM, [
        '**implicarea tinerilor:** echipa de cercetare include cel puțin un student cu vârsta de cel mult 24 de ani, angajat pe activități de cercetare-dezvoltare;',
        '**panelul de validare TRL 5** include **≥ 40 % femei**, **≥ 3 persoane peste 55 de ani** și **≥ 2 persoane cu deficiențe de vedere sau motorii**, recrutate inclusiv printr-o asociație de profil;',
        '**testul de accesibilitate** al interfeței SOUL (control vocal, modul „taste mari”, contrast ridicat, feedback haptic/sonor, dimensiunea textului), realizat împreună cu o asociație a persoanelor cu dizabilități, cu raport public;',
        '**accesibilitatea rezultatelor:** produsul rezultat nu depinde exclusiv de gesturi fine sau de vedere; documentația și materialele de comunicare sunt publicate în formate accesibile.',
    ]),
    ('5. Monitorizare', [
        'responsabilul desemnat ține evidența anonimizată a componenței echipei și a panelului de testare și raportează anual administratorului;',
        'politica se revizuiește anual sau la orice modificare legislativă relevantă.',
    ]),
]
for h, items in sections:
    d.add_heading(h, level=2)
    for it in items:
        P(d, it, 'j', bullet=True)
d.add_paragraph()
table(d, [], [[f'Administrator,\n[[Nume Prenume]]\n\nSemnătura: ____________________\nData: [[ZZ.LL.AAAA]]']], [16.6], head=False)
d.save(f'{OUTDIR}/Politica_egalitate_sanse_DECIZIE.docx')
print('ok')
