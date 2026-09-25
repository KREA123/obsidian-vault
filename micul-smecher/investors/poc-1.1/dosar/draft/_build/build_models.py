# -*- coding: utf-8 -*-
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docxlib import write_runs, W, set_rpr
from data import TITLU, APPLICANT

OUTDIR = sys.argv[1]
NAVY = RGBColor(0x1B, 0x3A, 0x6B)


def base_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(3)
    st.element.rPr.rFonts.set(W('eastAsia'), 'Calibri')
    for lvl, size in (('Heading 1', 13), ('Heading 2', 11)):
        h = d.styles[lvl]
        h.font.name = 'Calibri'; h.paragraph_format.space_before = Pt(6); h.paragraph_format.space_after = Pt(2); h.font.size = Pt(size); h.font.color.rgb = NAVY; h.font.bold = True
        rpr = h.element.get_or_add_rPr()
        rf = rpr.find(W('rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            rf.set(W(a), 'Calibri')
        for a in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            if rf.get(W(a)) is not None:
                del rf.attrib[W(a)]
    z = d.settings.element.find(W('zoom'))
    if z is not None:
        z.set(W('percent'), '100')
    return d


def P(d, text, align=None, size=None, italic=False, bullet=False):
    p = d.add_paragraph(style='List Bullet' if bullet else None)
    write_runs(p._p, text, plain=False, size=int(size * 2) if size else None)
    if italic:
        for r in p.runs:
            r.italic = True
    if align == 'c':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'r':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == 'j':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(W('val'), 'clear'); sh.set(W('color'), 'auto'); sh.set(W('fill'), fill)
    tcpr.append(sh)


def table(d, header, data, widths):
    t = d.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        write_runs(c.paragraphs[0]._p, f'**{h}**', plain=False, color='FFFFFF', size=19)
        shade(c, '1B3A6B')
    for row in data:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            lines = v.split('\n')
            write_runs(cells[j].paragraphs[0]._p, lines[0], plain=False, size=19)
            for ln in lines[1:]:
                write_runs(cells[j].add_paragraph()._p, ln, plain=False, size=19)
    t.autofit = False
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trpr.append(cs)
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    d.add_paragraph()
    return t


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def note_box(d, text, fill='FFF4E5'):
    t = d.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    c = t.rows[0].cells[0]
    shade(c, fill)
    lines = text.split('\n')
    write_runs(c.paragraphs[0]._p, lines[0], plain=False, size=19)
    for ln in lines[1:]:
        p = c.add_paragraph(); write_runs(p._p, ln, plain=False, size=19)
    d.add_paragraph()


# =========================================================== LETTERS OF INTENT
d = base_doc()
P(d, '**MODEL — SCRISOARE DE INTENȚIE / DE INTERES**', 'c', 14)
P(d, 'Proiectul SOUL-PoC · Apelul PR BI P1/1.1/1/2026 — Proof of Concept · Grila ETF, subcriteriul 4.2 (4 puncte)', 'c', 10, italic=True)
note_box(d, '**Cum se folosește acest model (se șterge înainte de tipărire):**\n'
         '1. Se tipărește **pe antetul emitentului** (nu al ARTEMIS) și se semnează de **reprezentantul legal** al emitentului (sau electronic, cu semnătură calificată).\n'
         '2. Se completează toate câmpurile [[galbene]]; **nu** se modifică titlul proiectului (trebuie să fie identic cu CF).\n'
         '3. Data trebuie să fie **cu cel mult 6 luni înainte de depunerea CF**; emitenții trebuie să fie **entități distincte** (KREA/MundiShop nu contează — sunt tot ARTEMIS).\n'
         '4. Pentru punctaj sunt necesare **≥ 2 scrisori valide**; recomandăm 3–4 (una de rezervă). O scrisoare căreia îi lipsește oricare dintre cele 7 elemente **nu se ia în calcul**.\n'
         '5. Scrisoarea exprimă o intenție și nu creează obligații financiare; nu se promite nicio plată către emitent.')
P(d, '**Verificarea celor 7 elemente obligatorii (Grila 4.2) — unde se află în fiecare variantă:**')
table(d, ['Element obligatoriu (Grila 4.2)', 'Unde apare în model'], [
    ['(i) Specific proiectului — referință la titlul proiectului sau la rezultatul vizat', 'secțiunea 1 „Proiectul vizat” (titlul complet + acronimul SOUL-PoC)'],
    ['(ii) Descrierea soluției tehnice/comerciale vizate', 'secțiunea 2 „Soluția vizată”'],
    ['(iii) Rolul partenerului (testare/pilot, co-dezvoltare, utilizare, distribuție, licențiere etc.)', 'secțiunea 3 „Rolul nostru”'],
    ['(iv) Forma de colaborare', 'secțiunea 4 „Forma colaborării”'],
    ['(v) Orizontul temporal (data de început și/sau durata)', 'secțiunea 5 „Orizontul de timp”'],
    ['(vi) Semnat de reprezentantul legal, cu nume, funcție, date de contact (e-mail, telefon)', 'blocul de semnătură'],
    ['(vii) Datată (≤ 6 luni înainte de depunerea CF)', 'antet („Nr. / Data”) + lângă semnătură'],
], [8.5, 8.1])

SOL = ('SOUL este un dispozitiv tangibil de interacțiune om–mașină / asistent personal AI dezvoltat de ARTEMIS DIGITAL S.R.L.: '
       'un obiect de ~90 × 103 × 31 mm, cu carcasă din aluminiu reciclabil și afișaj rotund de 2,8″, care arată doi ochi expresivi și '
       'are o tastatură circulară. Funcționează offline (notițe, alarme, mementouri, liste) sau cu asistentul AI ales de utilizator '
       '(Claude, ChatGPT sau cheie proprie), inclusiv pentru aprobarea cererilor Claude Code; datele sunt procesate pe servere din UE, '
       'iar bateria poate fi înlocuită de utilizator. Proiectul validează SOUL de la TRL 3 la TRL 5: 15 prototipuri pre-serie testate '
       'de ≥ 12 utilizatori din București-Ilfov timp de 4 săptămâni.')

variants = [
    ('A', 'retailer / distribuitor', 'magazin de electronice/obiecte de design premium, concept store, distribuitor online',
     ['evaluarea produsului SOUL pentru portofoliul nostru și furnizarea de feedback comercial structurat (preț perceput, poziționare la raft/online, ambalaj, garanție) pe baza prototipurilor validate la TRL 5;',
      'prezentarea a [[1–2]] prototipuri în [[magazinul/showroom-ul nostru din …]] în perioada de validare, pentru reacția clienților [[opțional]];',
      'după validarea TRL 5 și certificarea CE: negocierea unui contract de **distribuție** [[nonexclusivă, în România / UE]], cu o comandă inițială orientativă de [[X]] unități [[opțional, fără caracter obligatoriu]].'],
     'acord de colaborare comercială nonexclusiv, fără obligații financiare pentru nicio parte în perioada proiectului; feedback prin interviuri și chestionare; contract de distribuție negociat separat după validare.',
     'de la data semnării contractului de finanțare de către ARTEMIS DIGITAL S.R.L. (estimat în [[luna/anul — ex. 2027]]), pe durata implementării (24 de luni) și [[12]] luni după finalizare.'),
    ('B', 'laborator universitar / organizație de cercetare', 'facultate, departament, institut național de cercetare-dezvoltare',
     ['acces la infrastructura laboratorului [[denumirea laboratorului — ex. camera anecoică / laborator de radiofrecvență / EMC]] pentru măsurători RF prin carcasa din aluminiu (protocolul P8) și, după caz, pre-testare EMC/RF;',
      'expertiză științifică în [[antene și propagare / sisteme embedded / interacțiune om–calculator]] pentru validarea tehnologiei la TRL 4–5 (co-dezvoltare, consultare metodologică);',
      'sprijin pentru identificarea unui student [[ETTI/ACS]] ≤ 24 ani implicat în activitățile de cercetare și, după caz, participarea la redactarea unei publicații open-access comune.'],
     'acord de colaborare în cercetare-dezvoltare; serviciile de laborator/CDI se contractează separat, la prețul pieței și cu respectarea procedurilor de achiziție aplicabile (nu se promit sume prin această scrisoare).',
     'de la data semnării contractului de finanțare (estimat [[luna/anul]]), în lunile [[L4–L22]] ale proiectului (durata totală 24 de luni).'),
    ('C', 'utilizator pilot din mediul de afaceri', 'firmă de software/IT care folosește asistenți AI (ex. Claude Code) în activitatea zilnică',
     ['participarea la **pilotul de validare TRL 5**: [[3–5]] angajați ai noștri (dezvoltatori care folosesc zilnic Claude Code/asistenți AI) vor folosi prototipurile SOUL timp de 4 săptămâni în condiții reale de lucru, cu telemetrie opt-in și cu respectarea GDPR;',
      'furnizarea de feedback structurat (chestionar SUS, interviuri, jurnal de utilizare) privind aprobările Claude Code, notițele, alarmele și funcționarea offline;',
      'evaluarea, după validare, a achiziției a [[X]] unități pentru echipa noastră sau ca beneficii/cadouri corporate [[opțional, fără caracter obligatoriu]].'],
     'acord de pilot fără costuri pentru nicio parte; prototipurile rămân proprietatea ARTEMIS și se returnează la final; confidențialitate reciprocă; eventualul contract de achiziție se negociază separat.',
     'de la data semnării contractului de finanțare (estimat [[luna/anul]]); pilotul are loc în lunile [[L16–L22]] ale proiectului, iar colaborarea durează [[24]] de luni + [[6]] luni după finalizare.'),
]
for k, (code, kind, who, roles, form, horizon) in enumerate(variants):
    page_break(d)
    P(d, f'**VARIANTA {code} — {kind.upper()}**  ({who})', size=10, italic=True)
    P(d, '[[ANTETUL EMITENTULUI — denumire, CUI, adresă, e-mail, telefon]]', 'c')
    P(d, '**SCRISOARE DE INTENȚIE**', 'c', 14)
    P(d, 'Nr. [[…]] / Data: [[ZZ.LL.AAAA]]', 'r')
    P(d, f'Către: **{APPLICANT}**, CUI [[…]], [[adresa]], București')
    d.add_heading('1. Proiectul vizat', level=2)
    P(d, f'Subsemnatul/a [[Nume Prenume]], în calitate de **reprezentant legal** ([[funcția]]) al [[denumirea entității]] (CUI [[…]]), confirm interesul entității noastre de a colabora cu {APPLICANT} în legătură cu proiectul **„{TITLU}”** (acronim **SOUL-PoC**), propus spre finanțare în cadrul apelului PR BI P1/1.1/1/2026 — Proof of Concept, Programul Regional București-Ilfov 2021–2027.', 'j')
    d.add_heading('2. Soluția vizată', level=2)
    P(d, SOL, 'j')
    d.add_heading('3. Rolul nostru', level=2)
    for r in roles:
        P(d, r, bullet=True)
    d.add_heading('4. Forma colaborării', level=2)
    P(d, form[0].upper() + form[1:], 'j')
    d.add_heading('5. Orizontul de timp', level=2)
    P(d, 'Colaborarea începe ' + horizon, 'j')
    P(d, 'Prezenta scrisoare exprimă intenția noastră de colaborare și nu creează obligații financiare pentru niciuna dintre părți. Condițiile concrete se stabilesc prin acorduri/contracte separate.', 'j', italic=True)
    P(d, '**Reprezentant legal:** Nume și prenume: [[…]] · Funcția: [[reprezentant legal / administrator / director general / rector / director]] · Entitatea: [[denumire, CUI]]')
    P(d, '**Date de contact:** E-mail: [[…]] · Telefon: [[…]] · **Data:** [[ZZ.LL.AAAA — cu cel mult 6 luni înainte de depunerea CF]]')
    P(d, '**Semnătura și ștampila** (sau semnătură electronică calificată): ______________________________')
d.save(f'{OUTDIR}/Scrisoare_de_intentie_MODEL.docx')

# =========================================================== RIGHTS ASSIGNMENT
d = base_doc()
note_box(d, '**⚠ DE REVIZUIT DE UN AVOCAT ÎNAINTE DE SEMNARE.** Model de lucru redactat pentru dosarul SOUL-PoC (ghid §5.2.1 II, §5.2.2, §7.4 poz. 5: dovada drepturilor asupra rezultatelor CD de la care pornește proiectul; costul dobândirii acestor drepturi nu este eligibil, de aceea cesiunea este **cu titlu gratuit**). Puncte de verificat de avocat: (1) regimul operelor create cu asistență AI (codul a fost scris de autor împreună cu un asistent AI) — cesiunea acoperă drepturile „în măsura în care există”; (2) cesiunea desenului/modelului nedepus și a dreptului de a solicita înregistrarea (Legea 129/1992; Reg. (CE) 6/2002); (3) semnul „SOUL” nu este marcă înregistrată — se cedează doar drepturile/interesele existente și dreptul de a depune cererea; (4) dacă Andu este asociat/administrator, eventuala necesitate a unei hotărâri AGA pentru acceptare; (5) forma scrisă și data certă (ex. dată certă la notar/avocat), anterioară depunerii CF.', 'FDECEA')
P(d, '**CONTRACT DE CESIUNE A DREPTURILOR DE PROPRIETATE INTELECTUALĂ**', 'c', 14)
P(d, '**(declarație-contract de cesiune, cu titlu gratuit)**', 'c')
P(d, 'Nr. [[…]] / Data: [[ZZ.LL.AAAA — anterioară depunerii cererii de finanțare]]', 'c')
d.add_heading('1. Părțile', level=2)
P(d, '**Cedentul:** [[Nume Prenume]] („Andu”), cetățean român, domiciliat în [[adresa]], CNP [[…]], act de identitate [[CI seria … nr. …]], în calitate de autor/creator, denumit în continuare „**Cedentul**”;', 'j')
P(d, f'**Cesionarul:** **{APPLICANT}**, cu sediul în [[adresa]], București, înregistrată la Registrul Comerțului sub nr. [[J40/…/…]], CUI [[…]], reprezentată legal de [[Nume Prenume]], în calitate de administrator, denumită în continuare „**Cesionarul**”.', 'j')
P(d, 'Temei legal: Legea nr. 8/1996 privind dreptul de autor și drepturile conexe (în special art. 39–45 și art. 73–81 privind programele pentru calculator); Legea nr. 129/1992 privind protecția desenelor și modelelor; Regulamentul (CE) nr. 6/2002 privind desenele sau modelele comunitare; Legea nr. 84/1998 privind mărcile; Codul civil.', 'j', 10, italic=True)
d.add_heading('2. Obiectul cesiunii', level=2)
P(d, '2.1. Cedentul cesionează Cesionarului, **exclusiv**, **cu titlu gratuit**, **pe întreaga durată de protecție** și **pentru teritoriul întregii lumi**, toate drepturile patrimoniale de proprietate intelectuală pe care le deține, în măsura în care există, asupra rezultatelor enumerate în **Anexa 1** („Rezultatele SOUL”), create până la data prezentului contract, și anume:', 'j')
for x in ['codul sursă și obiect al firmware-ului „Suflet”/SoulOS, al serviciului AI `suflet_ai`, al simulatorului, al prototipului web SoulOS și al instrumentelor asociate (programe pentru calculator), inclusiv documentația;',
          'designul industrial al produsului SOUL (corpul SOUL, variantele v6/v7/S/M, capsula de încărcare „OU”, interfața grafică a „ochilor”), modelele CAD (CadQuery/OpenSCAD, STL/STEP), randările și desenele, inclusiv dreptul de a solicita înregistrarea ca desen/model industrial în orice țară și la EUIPO și dreptul de prioritate aferent;',
          'soluțiile tehnice, know-how-ul și rezultatele de cercetare (studii, specificații, protocoale, date de test), inclusiv dreptul de a solicita brevete sau modele de utilitate;',
          'semnul/denumirea „SOUL” și elementele de identitate vizuală (logo, nume de produs), inclusiv dreptul de a depune cereri de înregistrare ca marcă, precum și orice nume de domeniu și cont asociat [[de enumerat în Anexa 1]].']:
    P(d, x, bullet=True)
P(d, '2.2. Drepturile cedate includ, fără limitare, dreptul de reproducere, distribuire, închiriere, import, comunicare publică, punere la dispoziția publicului, traducere, adaptare, modificare și realizare de opere derivate, dreptul de a exploata comercial Rezultatele SOUL, de a acorda licențe (inclusiv open-source) și de a cesiona mai departe drepturile, prin orice mijloc și în orice formă, cunoscută sau necunoscută la data semnării.', 'j')
P(d, '2.3. Cesiunea **nu** cuprinde componentele terțe enumerate în **Anexa 2**, care rămân sub licențele lor (MIT, BSD, SIL OFL etc.); Cesionarul le va folosi cu respectarea acestor licențe.', 'j')
d.add_heading('3. Prețul', level=2)
P(d, '3.1. Cesiunea se face **cu titlu gratuit**. Cedentul declară că nu pretinde și nu va pretinde nicio remunerație sau despăgubire pentru drepturile cedate, inclusiv din viitoarea exploatare a Rezultatelor SOUL [[avocat: verificați renunțarea la remunerația proporțională/„bestseller” — art. 43 alin. (3) Legea 8/1996]].', 'j')
P(d, '3.2. Părțile iau act că, potrivit ghidului apelului PR BI P1/1.1/1/2026, costul dobândirii drepturilor asupra rezultatelor CD anterioare nu este eligibil și nu va fi solicitat la finanțare.', 'j')
d.add_heading('4. Declarațiile Cedentului', level=2)
for x in ['este autorul/creatorul Rezultatelor SOUL (singur, cu asistența unor instrumente software, inclusiv asistenți AI), iar datele creației sunt atestate de istoricul depozitului git (primul commit: 24.09.2026; commit de referință: [[hash la data semnării]]);',
          'nu a cedat și nu a licențiat anterior, exclusiv sau neexclusiv, drepturile asupra Rezultatelor SOUL către terți și nu există litigii sau revendicări privind aceste drepturi;',
          'Rezultatele SOUL nu au fost create în executarea unui contract individual de muncă sau a unui alt contract cu un terț care ar conferi drepturi acestuia [[avocat/Ct: de verificat, inclusiv raportul cu ARTEMIS dacă Andu este salariat]];',
          'toate componentele terțe cunoscute sunt enumerate în Anexa 2, cu licențele lor;',
          'va semna orice document suplimentar necesar pentru înregistrarea drepturilor pe numele Cesionarului (EUIPO, OSIM, registre de domenii) și va preda Cesionarului, la cerere, toate fișierele sursă.']:
    P(d, x, bullet=True)
d.add_heading('5. Drepturile morale', level=2)
P(d, '5.1. Drepturile morale de autor rămân la Cedent, potrivit legii. Cedentul este de acord ca Rezultatele SOUL să fie modificate, adaptate și comercializate de Cesionar fără mențiunea numelui său, cu excepția documentelor științifice (publicații, rapoarte), în care va fi menționat ca autor/co-autor [[avocat: formularea exactă]].', 'j')
d.add_heading('6. Dispoziții finale', level=2)
P(d, '6.1. Contractul intră în vigoare la data semnării de către ambele părți. 6.2. Orice modificare se face prin act adițional scris. 6.3. Litigiile se soluționează pe cale amiabilă, iar în caz contrar de instanțele competente din București. 6.4. Contractul s-a încheiat în [[2]] exemplare originale, câte unul pentru fiecare parte [[sau: semnat electronic cu semnătură calificată]].', 'j')
table(d, ['CEDENT', 'CESIONAR'], [
    ['[[Nume Prenume]]\n\nSemnătura: ____________________\nData: [[ZZ.LL.AAAA]]',
     f'{APPLICANT}\nprin [[Nume Prenume]], administrator\nSemnătura și ștampila: ____________________\nData: [[ZZ.LL.AAAA]]'],
], [8.3, 8.3])
page_break(d)
d.add_heading('Anexa 1 — Rezultatele SOUL cedate (inventar)', level=2)
table(d, ['Nr.', 'Rezultat', 'Locație / identificare', 'Data / versiune'], [
    ['1', 'Firmware „Suflet” / SoulOS (C++17): bibliotecile `lib/Suflet`, `src/`, `sim/`, `test/`, `tools/`', 'depozit git, `micul-smecher/firmware/`', '[[hash commit]] / 25.09.2026'],
    ['2', 'Serviciul AI `suflet_ai` (Python) + teste + exemple', '`micul-smecher/ai/`', '[[hash commit]]'],
    ['3', 'Prototipul web și specificația SoulOS', '`micul-smecher/os/` (index.html, SPEC.md, ARCHITECTURE.md, research/)', '[[hash commit]]'],
    ['4', 'Modele CAD, STL/STEP, prototip P0', '`micul-smecher/cad/`, `micul-smecher/prototip/`', '[[hash commit]]'],
    ['5', 'Randări și design industrial (v6/v7/S/M, capsula OU, „ochii”)', '`micul-smecher/renders/`, `media/`', '[[hash commit]]'],
    ['6', 'Identitate vizuală și denumirea „SOUL” (logo, nume)', '`micul-smecher/brand/`, `site/`', '[[hash commit]]'],
    ['7', 'Studii, specificații, protocoale, date', '`micul-smecher/research/`, `docs/`, `blueprints/`', '[[hash commit]]'],
    ['8', 'Nume de domeniu / conturi', '[[de completat]]', '[[…]]'],
], [1.0, 6.2, 6.0, 3.4])
d.add_heading('Anexa 2 — Componente terțe excluse din cesiune (se folosesc sub licențele lor)', level=2)
table(d, ['Componentă', 'Unde e folosită', 'Licență'], [
    ['anthropics/claude-desktop-buddy (Anthropic, PBC)', '`firmware/src/ble_link.cpp` (adaptat)', 'MIT'],
    ['78/xiaozhi-esp32 (Shenzhen Xinzhi Future Technology)', 'valori de registre PMU, hartă de pini; poarta vocală (planificat)', 'MIT'],
    ['Inigo Quilez — funcții SDF 2D', '`Canvas.cpp` (sdHeart, sdStar5)', 'MIT'],
    ['Fontul Nunito (The Nunito Project Authors)', '`FontData.h` (glife rasterizate)', 'SIL OFL 1.1'],
    ['GFX Library for Arduino; SensorLib; ArduinoJson; Unity', 'biblioteci PlatformIO (nevendorizate)', 'BSD; MIT; MIT; MIT'],
    ['Dependențele Python din `ai/requirements.txt` (ex. SDK Anthropic)', 'serviciul AI', '[[de listat cu licențele]]'],
], [6.2, 6.4, 4.0])
d.save(f'{OUTDIR}/Declaratie_cesiune_drepturi_MODEL.docx')
print('ok')
